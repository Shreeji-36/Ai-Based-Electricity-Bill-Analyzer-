import csv
import io
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

MACHINE_HEAD = ["Machine", "Qty", "Hours/day", "Units (kWh)", "Cost (INR)", "Share (%)", "Efficiency (%)"]
REC_HEAD = ["Machine", "Type", "Recommendation", "Units saved (kWh/month)", "Cost saved (INR/month)"]


def _machine_row(x: dict) -> list:
    return [x["name"], x.get("qty") or "", x.get("hours") or "", x["units"], x["cost"], x["pct"], x["efficiency"]]


def _recs(r: dict) -> list[dict]:
    """Prefer the machine-specific AI recommendations; fall back to the basic tips."""
    ai = (r.get("ai") or {}).get("recommendations")
    if ai:
        return ai
    return [{"machine": "", "kind": "", **t} for t in r.get("tips", [])]


def _total_savings(r: dict) -> int:
    ai = r.get("ai") or {}
    return ai.get("total_savings", r.get("total_savings", 0))


def _warning(r: dict) -> str:
    e = r.get("estimate") or {}
    if not e.get("scaled"):
        return ""
    return (f"Your machine quantities and hours add up to about {e['raw_units']:,} kWh, but the bill is "
            f"{r['bill']['units']:,.0f} kWh. Machine figures were scaled down to fit the bill; "
            "please re-check quantities, hours per day and bill units.")


def to_csv(r: dict) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(MACHINE_HEAD)
    for x in r["rows"]:
        w.writerow(_machine_row(x))
    w.writerow([])
    w.writerow(REC_HEAD)
    for t in _recs(r):
        w.writerow([t.get("machine", ""), t.get("kind", ""), t["text"], t["save_units"], t["save_cost"]])
    return buf.getvalue().encode("utf-8-sig")  # BOM so Excel opens it correctly


def to_xlsx(r: dict) -> bytes:
    wb = Workbook()
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0F766E")

    def style_header(ws, widths):
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for col, wd in zip("ABCDEFG", widths):
            ws.column_dimensions[col].width = wd
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Machine Analysis"
    ws.append(MACHINE_HEAD)
    for x in r["rows"]:
        ws.append(_machine_row(x))
    style_header(ws, [30, 8, 11, 14, 14, 11, 14])

    rec = wb.create_sheet("Recommendations")
    rec.append(REC_HEAD)
    for t in _recs(r):
        rec.append([t.get("machine", ""), t.get("kind", ""), t["text"], t["save_units"], t["save_cost"]])
    style_header(rec, [26, 18, 80, 22, 22])
    for row in rec.iter_rows(min_row=2):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")

    summ = wb.create_sheet("Summary")
    b = r["bill"]
    ai = r.get("ai") or {}
    summ.append(["Item", "Value"])
    for k, v in [
        ("Industry", r["industry"]),
        ("Total units (kWh)", b["units"]), ("Total amount (INR)", b["amount"]),
        ("Billing days", b["days"]), ("Rate (INR/unit)", r["rate"]),
        ("Forecast next month (kWh)", (ai.get("forecast", {}).get("units") or [r["forecast"]["units"]])[0]),
        ("Forecast next month (INR)", (ai.get("forecast", {}).get("amount") or [r["forecast"]["amount"]])[0]),
        ("Estimated savings (INR/month)", _total_savings(r)),
        ("Bill check", (ai.get("bill_anomaly") or {}).get("message", "")),
        ("Input warning", _warning(r)),
    ]:
        summ.append([k, v])
    style_header(summ, [32, 90])
    for row in summ.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(r: dict) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    s = getSampleStyleSheet()
    small = s["BodyText"].clone("small", fontSize=8.5, leading=11)
    head_style = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F766E")),
                  ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                  ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                  ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
                  ("VALIGN", (0, 0), (-1, -1), "TOP"),
                  ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")])]

    b, ai = r["bill"], r.get("ai") or {}
    fc = ai.get("forecast") or {}
    nxt_u = (fc.get("units") or [r["forecast"]["units"]])[0]
    nxt_a = (fc.get("amount") or [r["forecast"]["amount"]])[0]

    items = [
        Paragraph(f"Energy Report - {escape(r['industry'])} Industry", s["Title"]),
        Paragraph(f"Units: {b['units']:,.0f} kWh | Amount: INR {b['amount']:,.0f} | "
                  f"Rate: INR {r['rate']}/unit | Billing days: {b['days']}", s["Normal"]),
        Spacer(1, 8),
    ]
    if _warning(r):
        items += [Paragraph(f"<b>Warning:</b> {escape(_warning(r))}", small), Spacer(1, 6)]
    if (ai.get("bill_anomaly") or {}).get("message"):
        items += [Paragraph(f"Bill check: {escape(ai['bill_anomaly']['message'])}", small), Spacer(1, 6)]

    rows = [["Machine", "Qty", "Hrs/day", "Units (kWh)", "Cost (INR)", "Share %"]]
    for x in r["rows"]:
        rows.append([Paragraph(escape(x["name"]), small), x.get("qty") or "-", x.get("hours") or "-",
                     f"{x['units']:,}", f"{x['cost']:,}", x["pct"]])
    t = Table(rows, colWidths=[170, 40, 50, 90, 90, 60], repeatRows=1)
    t.setStyle(TableStyle(head_style + [("ALIGN", (1, 0), (-1, -1), "RIGHT")]))
    items += [Paragraph("Machine-wise Analysis", s["Heading2"]), t, Spacer(1, 10),
              Paragraph(f"Forecast next month: {nxt_u:,} units / INR {nxt_a:,}", s["Normal"]),
              Spacer(1, 6), Paragraph("Recommendations", s["Heading2"])]

    recs = _recs(r)
    if recs:
        rt = [["Machine", "Recommendation", "Saves (INR/mo)"]]
        for x in recs:
            kind = f" <i>({escape(x['kind'])})</i>" if x.get("kind") else ""
            rt.append([Paragraph(escape(x.get("machine") or "-"), small),
                       Paragraph(escape(x["text"]) + kind, small), f"{x['save_cost']:,}"])
        tt = Table(rt, colWidths=[110, 330, 70], repeatRows=1)
        tt.setStyle(TableStyle(head_style + [("ALIGN", (2, 0), (2, -1), "RIGHT")]))
        items.append(tt)
    else:
        items.append(Paragraph("No major issues found. Keep monitoring monthly.", s["Normal"]))
    items += [Spacer(1, 8), Paragraph(f"Total estimated savings: INR {_total_savings(r):,} per month", s["Heading3"]),
              Paragraph("Estimates are based on rated machine power, the quantities and hours entered, "
                        "and typical savings ranges. Actual savings depend on site conditions.", small)]
    doc.build(items)
    return buf.getvalue()